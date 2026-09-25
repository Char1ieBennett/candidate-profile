#!/usr/bin/env python3
"""Render a candidate profile (profile.json) as a .docx, optionally a .pdf.

Usage:
    python build_profile.py profile.json --out Profile_NewYork_Staff_Software_Engineer.docx [--pdf]

Input schema (see ../references/profile-rules.md):
    {
      "title": "STAFF SOFTWARE ENGINEER",
      "specialisms": ["Platform Services", "Distributed Systems", "Healthcare Data"],
      "metro": "New York metro area",
      "years": 12,
      "summary": "…",
      "core_skills": ["…", "…"],
      "roles": [{"title": "…", "start": 2020, "end": "Present", "employer_type": "…", "bullets": ["…"]}],
      "education": ["Bachelor's degree in Computer Science"],
      "certifications": [],            # optional; section omitted when empty
      "systems": ["…", "…"]
    }

The document contains only the resume sections. After saving, the script re-reads the
document text and exits non-zero if any banned token (labels, placeholders, contact details)
is present, so nothing beyond the resume can reach the output.

Only dependency: python-docx. PDF conversion uses LibreOffice (`soffice`) when it is on PATH.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Arial"
GREY = RGBColor(0x66, 0x66, 0x66)
SEP = "   |   "

# Anything in this list must never appear in the finished document.
BANNED = [
    # label / disclaimer language
    r"representative (profile|candidate|resume)", r"illustrative", r"not a specific", r"dummy",
    r"placeholder", r"fictional", r"anonymi[sz]ed", r"this (profile|document|resume) (is|represents|describes)",
    r"how to use", r"delete this",
    # unfilled template markers
    r"\[", r"\]", r"\{", r"\}", r"xxx", r"tbd", r"lorem",
    # contact details / identity
    r"@", r"https?://", r"www\.", r"linkedin\.com", r"\(\d{3}\)\s?\d{3}", r"\b\d{3}[-.]\d{3}[-.]\d{4}\b",
    # month-level dates
    r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{4}\b",
]


# ---------- layout helpers ----------

def _fmt(p, size=11, before=0, after=6, line=1.25, keep_next=False):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    for r in p.runs:
        r.font.name = FONT
        r.font.size = Pt(size)


def _rule(p):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "2"); b.set(qn("w:color"), "999999")
    pbdr.append(b); pPr.append(pbdr)


def para(doc, text, size=11, bold=False, italic=False, color=None, before=0, after=6, line=1.25, keep_next=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if color is not None:
        r.font.color.rgb = color
    _fmt(p, size=size, before=before, after=after, line=line, keep_next=keep_next)
    return p


def section(doc, title):
    p = para(doc, title.upper(), size=11, bold=True, before=14, after=4, keep_next=True)
    _rule(p)
    return p


def bullets(doc, items, after_last=4):
    for k, it in enumerate(items):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(it)
        p.paragraph_format.left_indent = Inches(0.3)
        _fmt(p, before=0, after=after_last if k == len(items) - 1 else 3, line=1.25)


def role(doc, title, dates, employer, items):
    p = doc.add_paragraph()
    r = p.add_run(title); r.bold = True
    p.add_run("    " + dates).font.color.rgb = GREY
    _fmt(p, before=8, after=1, line=1.2, keep_next=True)
    para(doc, employer, italic=True, color=GREY, after=3, keep_next=True)
    bullets(doc, items, after_last=6)


def header(doc, title, subline, meta):
    para(doc, title, size=18, bold=True, after=2, line=1.1, keep_next=True)
    para(doc, subline, size=11, bold=True, after=2, keep_next=True)
    para(doc, meta, size=10.5, color=GREY, after=8, keep_next=True)


# ---------- render ----------

def _dates(r):
    end = r.get("end", "Present")
    return f"{r['start']} – {end}"


def render(doc, p: dict):
    years = int(p["years"])
    start = min(int(r["start"]) for r in p["roles"])
    header(doc, p["title"].upper(),
           "  /  ".join(p["specialisms"]),
           f"{p['metro']}{SEP}{years} years' experience{SEP}{start} – Present")

    section(doc, "Professional summary")
    para(doc, p["summary"], line=1.3, after=4)

    section(doc, "Core skills")
    para(doc, SEP.join(p["core_skills"]), line=1.3, after=4)

    section(doc, "Experience")
    for r in p["roles"]:
        role(doc, r["title"], _dates(r), r["employer_type"], r["bullets"])

    section(doc, "Education")
    bullets(doc, p["education"])

    if p.get("certifications"):
        section(doc, "Certifications / licenses")
        bullets(doc, p["certifications"])

    section(doc, "Systems")
    para(doc, SEP.join(p["systems"]), line=1.3, after=4)


def build(profile: dict, out: Path) -> Path:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = FONT; normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    render(doc, profile)

    # No author / no stray identity in file properties.
    cp = doc.core_properties
    cp.author = ""; cp.last_modified_by = ""; cp.comments = ""; cp.subject = ""; cp.keywords = ""
    cp.title = profile["title"].title()

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


# ---------- checks ----------

def docx_text(path: Path) -> str:
    d = Document(path)
    return "\n".join(par.text for par in d.paragraphs)


def check(profile: dict, path: Path) -> list[str]:
    text = docx_text(path)
    low = text.lower()
    problems = [f"banned token /{pat}/" for pat in BANNED if re.search(pat, low)]

    years = int(profile["years"])
    roles = profile["roles"]
    if str(years) not in profile["summary"]:
        problems.append(f"summary does not state {years} years")
    if profile["metro"] not in text:
        problems.append("metro missing from header")
    if not profile["metro"].endswith("metro area"):
        problems.append("metro must end with 'metro area'")
    if re.search(r",\s*(?!DC\b)[A-Z]{2}\b", profile["metro"]):  # "Washington, DC" is the metro's name
        problems.append("metro carries a state code")
    titles = [r["title"].strip().lower() for r in roles]
    if len(set(titles)) != len(titles):
        problems.append("role titles repeat — progression required")
    if roles and roles[0].get("end", "Present") != "Present":
        problems.append("first role must be the current one (end = Present)")
    for r in roles:
        if r.get("end", "Present") != "Present" and int(r["end"]) <= int(r["start"]):
            problems.append(f"role '{r['title']}' has end <= start")
    if not (1 <= len(roles) <= 4):
        problems.append("expected 1–4 roles")
    return problems


def to_pdf(docx_path: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("PDF skipped: LibreOffice (soffice) not on PATH — export the .docx to PDF yourself.")
        return None
    pdf = docx_path.with_suffix(".pdf")
    pdf.unlink(missing_ok=True)
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                       check=True, capture_output=True, timeout=180)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as e:
        print(f"PDF skipped: LibreOffice failed ({e.__class__.__name__}) — export the .docx to PDF yourself.")
        return None
    if not pdf.exists():
        print("PDF skipped: LibreOffice produced no file — export the .docx to PDF yourself.")
        return None
    return pdf


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("profile", help="profile.json")
    ap.add_argument("--out", required=True,
                    help="output .docx (Profile_<Metro>_<Role>.docx); a bare file name is saved to ~/Downloads")
    ap.add_argument("--pdf", action="store_true", help="also convert to PDF with LibreOffice if available")
    a = ap.parse_args()

    profile = json.loads(Path(a.profile).read_text())
    out = Path(a.out).expanduser()
    if out.parent == Path("."):  # bare file name → ~/Downloads (falls back to the cwd if there is none)
        downloads = Path.home() / "Downloads"
        out = (downloads if downloads.is_dir() else Path.cwd()) / out
    out = out.resolve()
    if out.suffix.lower() != ".docx":
        out = out.with_suffix(".docx")

    build(profile, out)
    problems = check(profile, out)
    if problems:
        out.unlink(missing_ok=True)
        print("REJECTED — fix profile.json and rebuild:", file=sys.stderr)
        for pr in problems:
            print(f"  - {pr}", file=sys.stderr)
        sys.exit(2)

    print(f"✓ {out}")
    if a.pdf:
        pdf = to_pdf(out)
        if pdf:
            print(f"✓ {pdf}")


if __name__ == "__main__":
    main()
